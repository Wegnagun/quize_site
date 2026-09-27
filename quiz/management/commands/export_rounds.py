import os
from datetime import datetime
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.template.loader import render_to_string
from weasyprint import HTML
from docx import Document
from docx.shared import Pt
from quiz.models import Quiz, Block, Question

class Command(BaseCommand):
    help = 'Exports current quiz Rounds and Questions to printable PDF and DOCX files'

    def add_arguments(self, parser):
        parser.add_argument(
            '--output',
            type=str,
            default='exports',
            help='Output directory name inside media/'
        )

    def handle(self, *args, **options):
        output_dir_name = options['output']
        
        try:
            # Получаем активный квиз
            quiz = Quiz.objects.last()
            if not quiz:
                raise CommandError("No Quiz found in database.")
            
            # Получаем блоки (раунды) этого квиза
            blocks = list(Block.objects.filter(quiz=quiz).order_by('order', 'title'))
            if not blocks:
                raise CommandError("Quiz has no rounds/blocks.")
                
        except Exception as e:
            raise CommandError(f"Database error: {e}")

        # --- СБОР ДАННЫХ ---
        block_ids = [b.id for b in blocks]
        
        # Загружаем все вопросы к этим блокам одним запросом
        all_questions = Question.objects.filter(block_id__in=block_ids).order_by('id')
        
        # Формируем надежный словарь: {ID_блока: [Объекты вопросов]}
        questions_map = {}
        for q in all_questions:
            if q.block_id not in questions_map:
                questions_map[q.block_id] = []
            questions_map[q.block_id].append(q)

        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
        base_path = os.path.join(settings.MEDIA_ROOT, output_dir_name, f"{quiz.title}_Rounds_{timestamp}")
        os.makedirs(base_path, exist_ok=True)

        self.stdout.write(self.style.SUCCESS(f"Exporting Round Sheets to: {base_path}"))

        context_pdf = {
            'quiz': quiz,
            'blocks': blocks,
            'questions_map': questions_map,
            'generated_at': datetime.now().strftime("%d.%m.%Y %H:%M"),
        }
        
        html_str_pdf = render_to_string('quiz_export/rounds_pdf.html', context_pdf)
        pdf_path_host = os.path.join(base_path, "ROUNDS_FOR_HOST.pdf")
        
        HTML(string=html_str_pdf).write_pdf(pdf_path_host)
        self.stdout.write(self.style.SUCCESS(f"Created Host Sheet: ROUNDS_FOR_HOST.pdf"))

        # --- ГЕНЕРАЦИЯ DOCX ---
        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

        heading = doc.add_heading(f'Ответы к Раундам: {quiz.title}', level=1)
        heading.alignment = 1
        
        p_date = doc.add_paragraph(f'Дата выгрузки: {datetime.now().strftime("%d.%m.%Y %H:%M")}')
        p_date.alignment = 1

        for block in blocks:
            doc.add_page_break() # Каждый раунд с новой страницы
            doc.add_heading(block.title, level=2)
            
            table = doc.add_table(rows=1, cols=2)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№ Вопроса'
            hdr_cells[1].text = 'Правильный ответ'
            
            block_q_list = questions_map.get(block.id, [])
            for idx, question in enumerate(block_q_list, start=1):
                row_cells = table.add_row().cells
                row_cells[0].text = str(idx)
                row_cells[1].text = question.answer or "—"
            
        docx_path = os.path.join(base_path, "ANSWERS_ROUNDS_REFERENCE.docx")
        doc.save(docx_path)
        self.stdout.write(self.style.SUCCESS(f"Created Answers Docx: ANSWERS_ROUNDS_REFERENCE.docx"))

        self.stdout.write(self.style.SUCCESS("\n--- EXPORT ROUNDS COMPLETE ---"))