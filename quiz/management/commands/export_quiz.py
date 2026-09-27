import os
from datetime import datetime
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.template.loader import render_to_string
from weasyprint import HTML
from docx import Document
from docx.shared import Pt
from quiz.models import Quiz, Task, Task_question # ВАЖНО: импортируем Task_question

class Command(BaseCommand):
    help = 'Exports the current quiz to printable PDF and DOCX files'

    def add_arguments(self, parser):
        parser.add_argument(
            '--output',
            type=str,
            default='exports',
            help='Output directory name inside media/'
        )

    def handle(self, *args, **options):
        output_dir_name = options['output']
        
        try:
            quiz = Quiz.objects.last()
            if not quiz:
                raise CommandError("No Quiz found in database.")
            
            # Получаем задачи ТЕКУЩЕГО квиза
            tasks = list(Task.objects.filter(quiz=quiz).order_by('title'))
            if not tasks:
                raise CommandError("Quiz has no Special Tasks.")
                
        except Exception as e:
            raise CommandError(f"Database error: {e}")

        # --- СБОР ДАННЫХ ---
        # Вместо одного сложного словаря сделаем простой поиск прямо в цикле шаблона 
        # ИЛИ соберем надежный dict вручную
        
        task_ids = [t.id for t in tasks]
        
        # Загружаем ВСЕ вопросы ко всем этим задачам одним запросом
        all_tq_objects = Task_question.objects.filter(task_id__in=task_ids).order_by('id')
        
        # Создаем четкий словарь: ключ - ID задачи, значение - список объектов вопросов
        tq_map = {}
        for obj in all_tq_objects:
            if obj.task_id not in tq_map:
                tq_map[obj.task_id] = []
            tq_map[obj.task_id].append(obj)

        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
        base_path = os.path.join(settings.MEDIA_ROOT, output_dir_name, f"{quiz.title}_{timestamp}")
        os.makedirs(base_path, exist_ok=True)

        self.stdout.write(self.style.SUCCESS(f"Exporting to: {base_path}"))

        context_pdf = {
            'quiz': quiz,
            'tasks': tasks,
            'tq_map': tq_map, # Передаем наш надежный словарь
            'generated_at': datetime.now().strftime("%d.%m.%Y %H:%M"),
        }
        
        html_str_pdf = render_to_string('quiz_export/questions_pdf.html', context_pdf)
        pdf_path_host = os.path.join(base_path, "QUIZ_FOR_HOST.pdf")
        
        HTML(string=html_str_pdf).write_pdf(pdf_path_host)
        self.stdout.write(self.style.SUCCESS(f"Created Host Sheet: QUIZ_FOR_HOST.pdf"))

        # --- ГЕНЕРАЦИЯ DOCX ---
        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

        heading = doc.add_heading(f'Ответы к квизу: {quiz.title}', level=1)
        heading.alignment = 1
        
        p_date = doc.add_paragraph(f'Дата выгрузки: {datetime.now().strftime("%d.%m.%Y %H:%M")}')
        p_date.alignment = 1

        for task in tasks:
            doc.add_page_break()
            doc.add_heading(task.title, level=2)
            
            if task.description:
                desc_p = doc.add_paragraph(f'Описание: {task.description}')
                desc_p.italic = True
            
            table = doc.add_table(rows=1, cols=2)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№ Вопроса'
            hdr_cells[1].text = 'Правильный ответ'
            
            # Берем вопросы из нашего надежного словаря
            task_q_list = tq_map.get(task.id, [])
            for idx, question in enumerate(task_q_list, start=1):
                row_cells = table.add_row().cells
                row_cells[0].text = str(idx)
                row_cells[1].text = question.answer or "—"
            
        docx_path = os.path.join(base_path, "ANSWERS_REFERENCE.docx")
        doc.save(docx_path)
        self.stdout.write(self.style.SUCCESS(f"Created Answers Docx: ANSWERS_REFERENCE.docx"))

        self.stdout.write(self.style.SUCCESS("\n--- EXPORT COMPLETE ---"))